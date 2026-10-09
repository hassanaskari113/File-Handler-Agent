"""All agent tools. Every path goes through sandbox.safe_path -> nothing outside WORK_DIR.

Tools 1-3 are the starter tools (hardened). Tools 4-13 come from the mind map.
"""

import base64
import csv
import functools
import html as htmllib
import io
import json
import os
import re
import shlex
import shutil
import subprocess
import sys

from langchain.tools import tool

from config import BASE_DIR, WORK_DIR
from sandbox import SandboxError, rel, safe_path

MAX_READ = 6000  # chars returned by read_file (keeps tokens low; use offset to page)
MAX_OUT = 3000  # chars of program output returned by execute_file


# ----------------------------------------------------------------- helpers
def _safe(fn):
    """Turn any exception into a plain 'Error: ...' string the agent can read."""

    @functools.wraps(fn)
    def wrapper(*a, **k):
        try:
            return fn(*a, **k)
        except SandboxError as e:
            return f"Error: {e}"
        except Exception as e:  # noqa: BLE001
            return f"Error: {type(e).__name__}: {e}"

    return wrapper


def _clip(text: str, n: int, tail: bool = False) -> str:
    if len(text) <= n:
        return text
    cut = len(text) - n
    return (
        f"[...{cut} chars cut...]\n{text[-n:]}"
        if tail
        else f"{text[:n]}\n[...{cut} chars cut...]"
    )


def _need_file(path: str) -> str:
    p = safe_path(path)
    if not os.path.exists(p):
        raise FileNotFoundError(f"'{rel(p)}' does not exist.")
    if os.path.isdir(p):
        raise IsADirectoryError(f"'{rel(p)}' is a folder, not a file.")
    return p


def _size(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024


def _docx_blocks(p: str):
    from docx import Document

    return Document(p).iter_inner_content()


def _load_text(p: str) -> str:
    """Text of txt-like, .docx and .pdf files (used by summary / conversion)."""
    ext = os.path.splitext(p)[1].lower()
    if ext == ".docx":
        from docx.table import Table

        out = []
        for b in _docx_blocks(p):
            if isinstance(b, Table):
                out += [" | ".join(c.text for c in r.cells) for r in b.rows]
            else:
                out.append(b.text)
        return "\n".join(out)
    if ext == ".pdf":
        from pypdf import PdfReader

        return "\n".join((pg.extract_text() or "") for pg in PdfReader(p).pages)
    with open(p, "r", encoding="utf-8") as f:
        return f.read()


# ---------------------------------------------- 1-3  starter tools (hardened)
@tool
@_safe
def list_files(path: str = ".") -> str:
    """List the files and folders inside a folder of the working directory."""
    p = safe_path(path)
    if not os.path.isdir(p):
        return f"Error: '{rel(p)}' is not a folder."
    entries = sorted(
        os.listdir(p), key=lambda n: (not os.path.isdir(os.path.join(p, n)), n.lower())
    )
    if not entries:
        return f"'{rel(p)}' is empty."
    lines = []
    for name in entries[:150]:
        full = os.path.join(p, name)
        lines.append(
            f"[folder] {name}"
            if os.path.isdir(full)
            else f"[file]   {name} ({_size(os.path.getsize(full))})"
        )
    if len(entries) > 150:
        lines.append(f"... and {len(entries) - 150} more")
    return "\n".join(lines)


@tool
@_safe
def read_file(path: str, offset: int = 0) -> str:
    """Read a text file. Long files are cut to ~6000 chars; pass offset to continue."""
    p = _need_file(path)
    try:
        with open(p, "r", encoding="utf-8") as f:
            text = f.read()
    except UnicodeDecodeError:
        return f"Error: '{rel(p)}' is not a UTF-8 text file (binary?)."
    offset = max(0, int(offset))
    chunk = text[offset : offset + MAX_READ]
    if offset + MAX_READ < len(text):
        chunk += f"\n[...truncated: showed chars {offset}-{offset + MAX_READ} of {len(text)}; call again with offset={offset + MAX_READ}]"
    return chunk if text else f"'{rel(p)}' is empty."


@tool
@_safe
def write_file(path: str, content: str) -> str:
    """Create or overwrite a text file with its COMPLETE content. Creates missing folders."""
    p = safe_path(path)
    if os.path.isdir(p):
        return f"Error: '{rel(p)}' is a folder."
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "w", encoding="utf-8", newline="\n") as f:
        f.write(content)
    return (
        f"Wrote {len(content.splitlines())} lines ({len(content)} chars) to '{rel(p)}'."
    )


# ------------------------------------------------------------ 4  execute_file
def _clean_env() -> dict:
    drop = ("KEY", "TOKEN", "SECRET", "PASSWORD", "CREDENTIAL")
    env = {k: v for k, v in os.environ.items() if not any(d in k.upper() for d in drop)}
    env.update(
        PYTHONIOENCODING="utf-8", PYTHONDONTWRITEBYTECODE="1", PYTHONUNBUFFERED="1"
    )
    return env


@tool
@_safe
def execute_file(path: str, arguments: str = "", timeout: int = 30) -> str:
    """Run a Python (.py) file in the sandbox and return its REAL exit code, stdout and stderr."""
    p = _need_file(path)
    if not p.lower().endswith(".py"):
        return "Error: only Python (.py) files can be executed (they run inside the folder sandbox)."
    cmd = [sys.executable, str(BASE_DIR / "sandbox_runner.py"), WORK_DIR, p]
    if arguments.strip():
        cmd += shlex.split(arguments, posix=(os.name != "nt"))
    timeout = min(max(int(timeout), 1), 60)
    try:
        r = subprocess.run(
            cmd,
            cwd=WORK_DIR,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=timeout,
            env=_clean_env(),
            stdin=subprocess.DEVNULL,
        )
    except subprocess.TimeoutExpired as e:

        def _t(x):
            return x.decode("utf-8", "replace") if isinstance(x, bytes) else (x or "")

        return (
            f"Error: timed out after {timeout}s (process killed).\n"
            f"--- stdout so far ---\n{_clip(_t(e.stdout), MAX_OUT)}"
        )
    out, err = r.stdout.strip(), r.stderr.strip()
    parts = [
        f"exit code: {r.returncode}",
        f"--- stdout ---\n{_clip(out, MAX_OUT) if out else '(empty)'}",
    ]
    if err:
        parts.append(f"--- stderr ---\n{_clip(err, MAX_OUT, tail=True)}")
    return "\n".join(parts)


# ------------------------------------------------------------- 5  delete_file
PENDING_DELETES: list[
    dict
] = []  # filled by the tool, resolved by the user in the GUI / terminal


@tool
@_safe
def delete_file(path: str) -> str:
    """Request deletion of a file or folder. ONLY when the user explicitly asked.
    It is NOT deleted until the user approves in the UI."""
    p = safe_path(path)
    if not os.path.exists(p):
        return f"Error: '{rel(p)}' does not exist."
    if os.path.normcase(p) == os.path.normcase(os.path.realpath(WORK_DIR)):
        return "Error: the working folder itself cannot be deleted."
    if not any(d["path"] == p for d in PENDING_DELETES):
        PENDING_DELETES.append(
            {"path": p, "rel": rel(p), "kind": "folder" if os.path.isdir(p) else "file"}
        )
    return (
        f"PENDING: '{rel(p)}' has NOT been deleted. The user must approve it in the UI. "
        "Tell the user to click Approve."
    )


def resolve_delete(item: dict, approve: bool) -> str:
    """Called by the GUI/terminal after the user's decision (never by the agent)."""
    if item in PENDING_DELETES:
        PENDING_DELETES.remove(item)
    if not approve:
        return f"Deletion of '{item['rel']}' was denied by the user."
    try:
        p = safe_path(item["path"])
        shutil.rmtree(p) if os.path.isdir(p) else os.remove(p)
        return f"'{item['rel']}' was deleted (approved by the user)."
    except Exception as e:  # noqa: BLE001
        return f"Could not delete '{item['rel']}': {e}"


# ----------------------------------------------------------- 6  create_folder
@tool
@_safe
def create_folder(path: str, files: dict | None = None) -> str:
    """Create a project folder. files = optional {"name.py": "content"} created inside it.
    If no files are given a README.md is added."""
    p = safe_path(path)
    os.makedirs(p, exist_ok=True)
    files = files or {"README.md": f"# {os.path.basename(p)}\n"}
    made = []
    for name, content in files.items():
        fp = safe_path(os.path.join(p, str(name)))
        os.makedirs(os.path.dirname(fp), exist_ok=True)
        with open(fp, "w", encoding="utf-8", newline="\n") as f:
            f.write(str(content))
        made.append(rel(fp))
    return f"Created folder '{rel(p)}' with: " + ", ".join(made)


# ---------------------------------------------------------- 7  convert_files
TEXT_EXT = {
    ".txt",
    ".md",
    ".log",
    ".py",
    ".js",
    ".html",
    ".css",
    ".json",
    ".csv",
    ".tsv",
    ".xml",
    ".yaml",
    ".yml",
    ".ini",
    ".cfg",
    ".sql",
    ".sh",
    ".bat",
    ".rst",
}
IMG_EXT = {
    ".png": "PNG",
    ".jpg": "JPEG",
    ".jpeg": "JPEG",
    ".bmp": "BMP",
    ".gif": "GIF",
    ".webp": "WEBP",
    ".tiff": "TIFF",
}


def _to_docx(text: str, out: str):
    from docx import Document

    d = Document()
    for line in text.splitlines():
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        d.add_heading(m.group(2), len(m.group(1))) if m else d.add_paragraph(line)
    d.save(out)


@tool
@_safe
def convert_files(path: str, new_extension: str) -> str:
    """Convert a file to another extension if possible (csv<->json, md->html, html->txt,
    docx->txt/md, txt/md->docx, png/jpg/webp/bmp/gif, or plain text<->text). Original is kept."""
    src = _need_file(path)
    base, old = os.path.splitext(src)
    old = old.lower()
    new = "." + new_extension.strip().lstrip(".").lower()
    if new == old:
        return f"Error: the file is already '{old}'."
    dst = base + new
    if os.path.exists(dst):
        return f"Error: '{rel(dst)}' already exists; not overwriting."

    def write(text):
        with open(dst, "w", encoding="utf-8", newline="\n") as f:
            f.write(text)

    if old == ".csv" and new == ".json":
        with open(src, newline="", encoding="utf-8") as f:
            write(json.dumps(list(csv.DictReader(f)), indent=2, ensure_ascii=False))
    elif old == ".json" and new == ".csv":
        with open(src, encoding="utf-8") as f:
            data = json.load(f)
        rows = data if isinstance(data, list) else [data]
        if not all(isinstance(r, dict) for r in rows):
            return "Error: JSON must be an object or a list of objects to become CSV."
        cols = list(dict.fromkeys(k for r in rows for k in r))
        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)
        write(buf.getvalue())
    elif old == ".md" and new == ".html":
        import markdown

        with open(src, encoding="utf-8") as f:
            write(markdown.markdown(f.read(), extensions=["fenced_code", "tables"]))
    elif old in (".html", ".htm") and new == ".txt":
        with open(src, encoding="utf-8") as f:
            t = re.sub(r"(?is)<(script|style).*?</\1>", "", f.read())
        write(htmllib.unescape(re.sub(r"<[^>]+>", "", t)).strip())
    elif old == ".docx" and new in (".txt", ".md"):
        write(_load_text(src))
    elif old in (".txt", ".md") and new == ".docx":
        with open(src, encoding="utf-8") as f:
            _to_docx(f.read(), dst)
    elif old in IMG_EXT and new in IMG_EXT:
        from PIL import Image

        img = Image.open(src)
        if IMG_EXT[new] == "JPEG" and img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        img.save(dst, IMG_EXT[new])
    elif old in TEXT_EXT and new in TEXT_EXT:
        with open(src, encoding="utf-8") as f:
            text = f.read()
        if new == ".json":
            json.loads(text)  # must be valid JSON, otherwise "not possible"
        write(text)
    else:
        return f"Error: converting '{old}' to '{new}' is not possible/supported."
    return f"Converted '{rel(src)}' -> '{rel(dst)}' (original kept)."


# ------------------------------------------------------------- 8  create_pdf
@tool
@_safe
def create_pdf(path: str, output: str = "") -> str:
    """Create a PDF from a Word (.docx) file (also works for .txt / .md). Simple formatting only."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )
    from xml.sax.saxutils import escape

    src = _need_file(path)
    ext = os.path.splitext(src)[1].lower()
    if ext not in (".docx", ".txt", ".md"):
        return "Error: create_pdf needs a .docx (or .txt / .md) file."
    dst = safe_path(output) if output.strip() else os.path.splitext(src)[0] + ".pdf"
    st = getSampleStyleSheet()
    story = []

    if ext == ".docx":
        from docx.table import Table as DTable

        for b in _docx_blocks(src):
            if isinstance(b, DTable):
                data = [
                    [Paragraph(escape(c.text), st["BodyText"]) for c in r.cells]
                    for r in b.rows
                ]
                if data:
                    t = Table(data, repeatRows=1)
                    t.setStyle(
                        TableStyle(
                            [
                                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ]
                        )
                    )
                    story += [t, Spacer(1, 10)]
                continue
            if not b.text.strip():
                story.append(Spacer(1, 6))
                continue
            name = (b.style.name or "") if b.style is not None else ""
            txt = "".join(
                ("<b>" if r.bold else "")
                + ("<i>" if r.italic else "")
                + escape(r.text)
                + ("</i>" if r.italic else "")
                + ("</b>" if r.bold else "")
                for r in b.runs
            )
            if name == "Title":
                style = st["Title"]
            elif name.startswith("Heading"):
                lvl = "".join(c for c in name if c.isdigit()) or "1"
                style = st[f"Heading{min(int(lvl), 3)}"]
            else:
                style = st["BodyText"]
            if "List" in name:
                txt = "• " + txt
            story.append(Paragraph(txt or escape(b.text), style))
    else:
        with open(src, encoding="utf-8") as f:
            for line in f.read().splitlines():
                m = re.match(r"^(#{1,3})\s+(.*)", line)
                if m:
                    story.append(
                        Paragraph(escape(m.group(2)), st[f"Heading{len(m.group(1))}"])
                    )
                elif line.strip():
                    story.append(Paragraph(escape(line), st["BodyText"]))
                else:
                    story.append(Spacer(1, 6))

    os.makedirs(os.path.dirname(dst), exist_ok=True)
    SimpleDocTemplate(dst, pagesize=A4, title=os.path.basename(src)).build(
        story or [Spacer(1, 1)]
    )
    return f"Created PDF '{rel(dst)}' from '{rel(src)}'."


# ---------------------------------------------------------- 9  create_summary
@tool
@_safe
def create_summary(path: str) -> str:
    """Summarize a file (txt/code/docx/pdf) with the model and save it as <name>_summary.md."""
    from langchain_core.messages import HumanMessage, SystemMessage
    from llm import get_llm
    from throttle import THROTTLE

    src = _need_file(path)
    text = _load_text(src)
    if not text.strip():
        return f"Error: '{rel(src)}' has no readable text."
    note = ""
    if len(text) > 9000:
        text, note = (
            text[:9000],
            "\n\n_(Only the first 9000 characters were summarized.)_",
        )
    reply = get_llm().invoke(
        [
            SystemMessage(
                content="Summarize the user's file in clear Markdown: one-line overview, then "
                "3-7 key bullet points. Max 200 words. Use only what is in the text."
            ),
            HumanMessage(content=text),
        ],
        config={"callbacks": [THROTTLE]},
    )
    summary = (
        reply.content if isinstance(reply.content, str) else str(reply.content)
    ).strip()
    if not summary:
        return "Error: the model returned an empty summary."
    base = os.path.splitext(src)[0] + "_summary"
    dst, i = base + ".md", 2
    while os.path.exists(dst):
        dst, i = f"{base}{i}.md", i + 1
    with open(dst, "w", encoding="utf-8", newline="\n") as f:
        f.write(summary + note + "\n")
    return f"Summary saved to '{rel(dst)}'.\n{_clip(summary, 700)}"


# -------------------------------------------------- 10/11  encrypt / decrypt
MAGIC = "ENCRYPTED-v1"


def _fernet(password: str, salt: bytes):
    from cryptography.fernet import Fernet
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC

    kdf = PBKDF2HMAC(
        algorithm=hashes.SHA256(), length=32, salt=salt, iterations=200_000
    )
    return Fernet(base64.urlsafe_b64encode(kdf.derive(password.encode("utf-8"))))


def _atomic_write(p: str, data: bytes):
    tmp = p + ".tmp"
    with open(tmp, "wb") as f:
        f.write(data)
    os.replace(tmp, p)


@tool
@_safe
def encrypt_file(path: str, password: str) -> str:
    """Encrypt a file with a password given by the user (AES via Fernet) and overwrite it
    with the encrypted text."""
    p = _need_file(path)
    if len(password) < 4:
        return "Error: password must be at least 4 characters (ask the user for one)."
    with open(p, "rb") as f:
        data = f.read()
    if data.startswith(MAGIC.encode()):
        return f"Error: '{rel(p)}' is already encrypted."
    salt = os.urandom(16)
    token = _fernet(password, salt).encrypt(data).decode()
    _atomic_write(p, f"{MAGIC}\n{base64.b64encode(salt).decode()}\n{token}\n".encode())
    return f"Encrypted '{rel(p)}' (overwritten with encrypted text)."


@tool
@_safe
def decrypt_file(path: str, password: str, overwrite: bool = True) -> str:
    """Decrypt a file made by encrypt_file and show the plain text. overwrite=True also
    restores the plain text into the file."""
    from cryptography.fernet import InvalidToken

    p = _need_file(path)
    with open(p, "r", encoding="utf-8", errors="replace") as f:
        lines = f.read().splitlines()
    if len(lines) < 3 or lines[0] != MAGIC:
        return f"Error: '{rel(p)}' is not an encrypted file made by encrypt_file."
    try:
        data = _fernet(password, base64.b64decode(lines[1])).decrypt(lines[2].encode())
    except InvalidToken:
        return "Error: wrong password (or the file is corrupted). File left unchanged."
    if overwrite:
        _atomic_write(p, data)
    try:
        shown = _clip(data.decode("utf-8"), 3000)
    except UnicodeDecodeError:
        shown = "(binary content, not shown)"
    state = "restored in the file" if overwrite else "file left encrypted"
    return f"Decrypted '{rel(p)}' ({state}). Plain text:\n{shown}"


# ------------------------------------------------------------ 12  upload_file
@tool
@_safe
def upload_file(path: str, drive_folder_id: str = "") -> str:
    """Upload a file from the working folder to the user's Google Drive."""
    p = _need_file(path)
    creds_path = BASE_DIR / os.getenv("GOOGLE_CREDENTIALS_FILE", "credentials.json")
    token_path = BASE_DIR / "token.json"
    if not creds_path.exists():
        return (
            "Error: Google Drive is not set up. The user must put an OAuth 'Desktop app' "
            "credentials.json in the project folder (see README, section Google Drive)."
        )
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload

    scopes = ["https://www.googleapis.com/auth/drive.file"]
    creds = (
        Credentials.from_authorized_user_file(str(token_path), scopes)
        if token_path.exists()
        else None
    )
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            creds = InstalledAppFlow.from_client_secrets_file(
                str(creds_path), scopes
            ).run_local_server(port=0)
        token_path.write_text(creds.to_json(), encoding="utf-8")
    service = build("drive", "v3", credentials=creds, cache_discovery=False)
    meta = {"name": os.path.basename(p)}
    if drive_folder_id.strip():
        meta["parents"] = [drive_folder_id.strip()]
    res = (
        service.files()
        .create(
            body=meta,
            media_body=MediaFileUpload(p, resumable=False),
            fields="id,webViewLink",
        )
        .execute()
    )
    return f"Uploaded '{rel(p)}' to Google Drive. id={res.get('id')} link={res.get('webViewLink')}"


# ------------------------------------------------------------ 13  rename_file
@tool
@_safe
def rename_file(path: str, new_name: str) -> str:
    """Rename a file or folder (new_name is just the new name, same location)."""
    p = safe_path(path)
    if not os.path.exists(p):
        return f"Error: '{rel(p)}' does not exist."
    if os.path.normcase(p) == os.path.normcase(os.path.realpath(WORK_DIR)):
        return "Error: the working folder itself cannot be renamed."
    if (
        not new_name.strip()
        or any(c in new_name for c in "/\\")
        or new_name in (".", "..")
    ):
        return "Error: new_name must be a plain file name without slashes."
    dst = safe_path(os.path.join(os.path.dirname(p), new_name.strip()))
    if os.path.exists(dst):
        return f"Error: '{rel(dst)}' already exists; not overwriting."
    os.rename(p, dst)
    return f"Renamed '{rel(p)}' -> '{rel(dst)}'."


TOOLS = [
    list_files,
    read_file,
    write_file,
    execute_file,
    delete_file,
    create_folder,
    convert_files,
    create_pdf,
    create_summary,
    encrypt_file,
    decrypt_file,
    upload_file,
    rename_file,
]
